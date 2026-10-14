from types import SimpleNamespace

from docx import Document
from docx.oxml.ns import qn

from app.services.docx_builder import build_report_docx


def _report(**overrides):
    data = dict(
        id=1,
        title="Лабораторна робота №1",
        discipline="Штучний інтелект",
        teacher="Іваненко І.І.",
        group="КНТ-22-1",
        goal="Навчитися щось робити.",
        sections=[
            {
                "title": "Хід роботи",
                "text": "Текст.",
                "code": "print(1)",
                "subsections": [{"title": "Налаштування середовища", "text": "Підрозділ."}],
            },
            {"title": "", "text": "Без назви."},
        ],
        conclusion="У ході виконання роботи було...",
    )
    data.update(overrides)
    return SimpleNamespace(**data)


def _build(tmp_path, monkeypatch, **overrides):
    monkeypatch.setenv("GENERATED_DIR", str(tmp_path))
    return Document(build_report_docx(_report(**overrides)))


def test_toc_has_entry_for_every_heading(tmp_path, monkeypatch):
    doc = _build(tmp_path, monkeypatch)
    toc = [p for p in doc.paragraphs if p.style.name.startswith("toc ")]
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading ")]

    assert headings == ["МЕТА РОБОТИ", "1 ХІД РОБОТИ", "1.1 Налаштування середовища", "2 РОЗДІЛ 2", "ВИСНОВКИ"]
    assert [p.style.name for p in toc] == ["toc 1", "toc 1", "toc 2", "toc 1", "toc 1"]
    assert len(toc) == len(headings)
    for entry, heading in zip(toc, headings):
        text = "".join(t.text for t in entry._p.iter(qn("w:t")))
        assert text.startswith(heading)
        assert text[len(heading):].isdigit()  # cached page number


def test_toc_field_wraps_all_entries(tmp_path, monkeypatch):
    doc = _build(tmp_path, monkeypatch)
    toc = [p for p in doc.paragraphs if p.style.name.startswith("toc ")]
    instr = [i.text for i in toc[0]._p.iter(qn("w:instrText"))]
    assert instr[0].strip().startswith("TOC")
    end_types = [f.get(qn("w:fldCharType")) for f in toc[-1]._p.iter(qn("w:fldChar"))]
    assert end_types[-2:] == ["end", "end"]  # PAGEREF end, then TOC end


def test_conclusion_in_document(tmp_path, monkeypatch):
    doc = _build(tmp_path, monkeypatch)
    texts = [p.text for p in doc.paragraphs]
    assert texts[texts.index("ВИСНОВКИ") + 1] == "У ході виконання роботи було..."


def test_a4_page(tmp_path, monkeypatch):
    section = _build(tmp_path, monkeypatch).sections[0]
    assert (round(section.page_width.mm), round(section.page_height.mm)) == (210, 297)


def test_apply_toc_pages_replaces_cached_numbers(tmp_path, monkeypatch):
    from app.services.docx_builder import apply_toc_pages

    monkeypatch.setenv("GENERATED_DIR", str(tmp_path))
    path = build_report_docx(_report())
    pages = {f"_Toc{i:08d}": 40 + i for i in range(1, 6)}
    apply_toc_pages(path, pages)

    doc = Document(path)
    toc = [p for p in doc.paragraphs if p.style.name.startswith("toc ")]
    numbers = ["".join(t.text for t in p._p.iter(qn("w:t"))).rsplit(" ", 1)[-1][-2:] for p in toc]
    assert numbers == ["41", "42", "43", "44", "45"]
    assert doc.settings.element.find(qn("w:updateFields")) is None
